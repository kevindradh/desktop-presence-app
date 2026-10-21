"""Exports match frozen database results and never publish partial/cancelled data."""
from dataclasses import replace
from datetime import timedelta
from pathlib import Path
from threading import Event
import os
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from app.exports.attendance import export_report
from app.models.entities import Member,Role
from app.services.auth import AuthError
from app.services.database import DataValidationError
from app.services.reports import ReportFilter
from app.services.tasks import TaskContext,TaskCancelled
from tests.test_auth import auth,add_role
from tests.test_database import store
from tests.test_attendance import attendance,scan,DAY
from tests.test_reports import reports

pytestmark=pytest.mark.database


def context():return TaskContext(Event(),lambda value:None)


def test_xlsx_matches_snapshot_and_treats_user_text_as_literal(attendance,store,tmp_path):
    data=attendance
    with store.db.transaction() as session:
        row=session.get(Member,data.member);row.member_code='=2+2';row.full_name='José Müller 東京 こんにちは'
    assert scan(data).saved
    service=reports(data,tmp_path);snapshot=service.build(data.token,ReportFilter(DAY,DAY))
    path=tmp_path/'report.xlsx'
    result=export_report(service,data.token,snapshot.id,path,'xlsx',context())
    assert result.rows==1 and path.is_file()
    book=load_workbook(path)
    assert book.sheetnames==['Kehadiran','Detail event']
    sheet=book['Kehadiran']
    assert sheet['B9'].value=='=2+2' and sheet['B9'].data_type=='s'
    assert sheet['C9'].value==snapshot.rows[0].name
    assert sheet['A9'].value.date()==DAY and sheet['G9'].value.hour==8
    assert sheet.freeze_panes=='D9' and len(sheet.tables)==1
    assert book['Detail event']['B9'].value=='=2+2'
    assert not list(tmp_path.glob('.stargate-report-*'))


def test_multi_page_pdf_and_all_xlsx_rows_unicode_long_names(attendance,store,tmp_path):
    from app.repositories.core import MemberRepository,ShiftRepository
    data=attendance
    assert scan(data).saved
    with store.db.transaction() as session:
        for i in range(70):
            member=MemberRepository(session).create(f'EXPORT-{i:03}',
                'José 東京 '+('Nama panjang untuk pemeriksaan pembungkusan teks '*3)+str(i),data.group,'synthetic/no-photo.png')
            ShiftRepository(session).assign(member.id,data.shift,DAY,DAY)
    service=reports(data,tmp_path);snapshot=service.build(data.token,ReportFilter(DAY,DAY))
    assert len(snapshot.rows)==71
    destination=Path(os.environ.get('STARGATE_REPORT_EVIDENCE',tmp_path))
    destination.mkdir(parents=True,exist_ok=True)
    xlsx=destination/'stage-12-synthetic.xlsx';pdf=destination/'stage-12-synthetic.pdf'
    export_report(service,data.token,snapshot.id,xlsx,'xlsx',context(),overwrite=True)
    export_report(service,data.token,snapshot.id,pdf,'pdf',context(),overwrite=True)
    sheet=load_workbook(xlsx)['Kehadiran']
    assert sheet.max_row==79 and sheet['C9'].value==snapshot.rows[0].name
    assert sheet.row_dimensions[10].height>32
    info=subprocess.run(['pdfinfo',str(pdf)],check=True,capture_output=True,text=True,encoding='utf-8').stdout
    pages=int(next(line.split(':')[1] for line in info.splitlines() if line.startswith('Pages:')))
    assert pages>1
    extracted=subprocess.run(['pdftotext','-enc','UTF-8',str(pdf),'-'],check=True,capture_output=True,text=True,encoding='utf-8').stdout
    assert 'José' in extracted and '東京' in extracted
    assert all(row.code in extracted for row in snapshot.rows)
    assert extracted.count('Halaman')==pages
    import re
    for page in extracted.split('\f'):
        if 'pembungkusan' in page:
            code=re.search(r'EXPORT-\d+|SG-TEST-001',page)
            assert code and code.start()<page.index('pembungkusan'), 'Normal rows must not begin with an orphan name continuation'


@pytest.mark.parametrize('kind',['xlsx','pdf'])
def test_empty_export_has_metadata_and_no_fabricated_rows(attendance,tmp_path,kind):
    data=attendance;service=reports(data,tmp_path)
    snapshot=service.build(data.token,ReportFilter(DAY,DAY,member='NONE'))
    path=tmp_path/('empty.'+kind)
    assert export_report(service,data.token,snapshot.id,path,kind,context()).rows==0
    if kind=='xlsx':
        assert load_workbook(path)['Kehadiran']['A9'].value=='Tidak ada data sesuai filter.'
    else:
        text=subprocess.run(['pdftotext','-enc','UTF-8',str(path),'-'],check=True,capture_output=True,text=True,encoding='utf-8').stdout
        assert 'Tidak ada data sesuai filter.' in text


@pytest.mark.parametrize('cause',['cancel','write','revoked','expired','exists','wrong_type','unsupported_unicode'])
def test_export_failures_leave_no_partial_file_and_preserve_existing(attendance,store,tmp_path,monkeypatch,cause):
    import app.exports.attendance as module
    data=attendance;service=reports(data,tmp_path)
    if cause=='unsupported_unicode':
        with store.db.transaction() as session:session.get(Member,data.member).full_name='Nama \U0001f642'
    snapshot=service.build(data.token,ReportFilter(DAY,DAY))
    kind='pdf' if cause=='unsupported_unicode' else 'xlsx'
    path=tmp_path/('fail.'+kind);ctx=context()
    real=module.write_xlsx
    if cause=='exists':path.write_text('ORIGINAL')
    if cause=='expired':data.mono[0]+=901
    if cause=='wrong_type':path=tmp_path/'wrong.csv'
    if cause in ('cancel','write','revoked'):
        def writing(snapshot,path,context):
            path.write_text('PARTIAL')
            if cause=='cancel':context.cancelled.set();context.checkpoint()
            if cause=='write':raise OSError('synthetic write failure')
            if cause=='revoked':data.auth.logout(data.token)
        monkeypatch.setattr(module,'write_xlsx',writing)
    expected=TaskCancelled if cause=='cancel' else AuthError if cause=='revoked' else DataValidationError
    with pytest.raises(expected):export_report(service,data.token,snapshot.id,path,kind,ctx)
    assert path.read_text()=='ORIGINAL' if cause=='exists' else not path.exists()
    assert not list(tmp_path.glob('.stargate-report-*'))


def test_cancel_during_publication_reports_success_not_false_cancellation(attendance,tmp_path,monkeypatch):
    import app.exports.attendance as module
    data=attendance;service=reports(data,tmp_path)
    snapshot=service.build(data.token,ReportFilter(DAY,DAY));ctx=context();path=tmp_path/'complete.xlsx'
    link=module.os.link
    def publish(source,target):
        link(source,target)
        ctx.cancelled.set()
    monkeypatch.setattr(module.os,'link',publish)
    assert export_report(service,data.token,snapshot.id,path,'xlsx',ctx).rows==1
    ctx.checkpoint()
    assert path.is_file() and ctx.committed.is_set()


def test_missing_destination_folder_is_safe_error(attendance,tmp_path):
    data=attendance;service=reports(data,tmp_path)
    snapshot=service.build(data.token,ReportFilter(DAY,DAY))
    with pytest.raises(DataValidationError):export_report(service,data.token,snapshot.id,tmp_path/'missing'/'a.xlsx','xlsx',context())
