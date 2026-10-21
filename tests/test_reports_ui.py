"""Viewer report filtering, complete export pagination, stable forms and lifecycle."""
import logging
import os
from pathlib import Path
from threading import Event

import pytest
from openpyxl import load_workbook
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QFileDialog

from app.services.errors import ErrorReporter
from app.settings.config import load_config
from app.ui.reports_demo import ReportsDemo
from app.ui.theme import apply_theme
from database.reports_demo import seed_reports_demo
from tests.test_auth import auth,NEW
from tests.test_database import store
from tests.test_session_ui import views,sign_in,activate
from tests.test_administration_ui import open_page
from tests.test_attendance import attendance,DAY

pytestmark=pytest.mark.database


@pytest.fixture
def demo(store,tmp_path,qt_app,wait_until):
    config=load_config(tmp_path,environ={});data=seed_reports_demo(store,config)
    logger=logging.Logger('reports-ui-test');logger.addHandler(logging.NullHandler())
    apply_theme(qt_app)
    view=ReportsDemo(config,logger,ErrorReporter(logger),data)
    view.show()
    wait_until(lambda:view.page.snapshot is not None and not view.page.busy,timeout=15000)
    yield view,data
    view.close()
    wait_until(lambda:not view.isVisible() and not view.page.busy,timeout=15000)
    view.deleteLater();qt_app.processEvents()


def test_viewer_demo_pagination_and_export_all_loaded_rows(demo,tmp_path,wait_until,monkeypatch):
    view,data=demo;page=view.page
    assert {route[0] for route in view.shell.routes}=={'reports'}
    assert len(page.snapshot.rows)==65 and page.table.rowCount()==50 and page.next.isEnabled()
    page.next.click();assert page.offset==50 and page.table.rowCount()==15
    path=tmp_path/'all-pages.xlsx'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:(str(path),'Excel'))
    page.xlsx_button.click()
    wait_until(lambda:not page.busy,timeout=15000)
    assert path.exists() and '65 baris' in page.message.text()
    assert load_workbook(path)['Kehadiran'].max_row==73
    assert page.offset==50 and page.table.rowCount()==15
    destination=os.environ.get('STARGATE_REPORT_EVIDENCE')
    if destination:
        Path(destination).mkdir(parents=True,exist_ok=True)
        assert view.grab().save(str(Path(destination)/'stage-12-reports.png'))
    page.previous.click();assert page.offset==0
    page.status_filter.setCurrentIndex(page.status_filter.findData('LATE'))
    assert page.snapshot is None and not page.xlsx_button.isEnabled()
    page.load_button.click()
    wait_until(lambda:not page.busy)
    assert len(page.snapshot.rows)==1 and page.snapshot.rows[0].status=='LATE'


def test_report_filter_focus_and_period_presets(attendance,views,wait_until):
    view=views(attendance.auth);view.reports_service.clock=lambda:attendance.now[0]
    sign_in(view,'root',NEW,wait_until)
    page=open_page(view,'reports',wait_until)
    page.period.setCurrentIndex(page.period.findData('WEEKLY'))
    assert page.start.date().dayOfWeek()==1 and page.start.date().daysTo(page.end.date())==6
    page.period.setCurrentIndex(page.period.findData('MONTHLY'))
    assert page.start.date().day()==1 and page.end.date().day()==31
    page.period.setCurrentIndex(page.period.findData('CUSTOM'));assert page.end.isEnabled()
    activate(view)
    page.member.setFocus();page.member.setText('Anggota Uji')
    for _ in range(3):
        view._check_session();QTest.qWait(80)
        assert page.member.hasFocus() and page.member.text()=='Anggota Uji'


@pytest.mark.parametrize('action',['cancel','logout','close'])
def test_export_cancel_and_leaving_do_not_publish_partial_file(demo,tmp_path,wait_until,monkeypatch,action):
    import app.exports.attendance as module
    view,data=demo;page=view.page
    began,release=Event(),Event();path=tmp_path/'cancelled.xlsx'
    def slow(snapshot,path,context):
        path.write_text('PARTIAL')
        began.set()
        assert release.wait(6)
        context.checkpoint()
    monkeypatch.setattr(module,'write_xlsx',slow)
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:(str(path),'Excel'))
    page.xlsx_button.click()
    try:
        wait_until(began.is_set)
        assert page.cancel_button.isEnabled() and not page.member.isEnabled()
        if action=='cancel':page.cancel_button.click()
        elif action=='logout':view.logout()
        else:view.close()
        release.set()
        wait_until(lambda:not page.busy,timeout=10000)
        if action=='logout':wait_until(lambda:view.login is None)
        if action=='close':wait_until(lambda:not view.isVisible(),timeout=10000)
        assert not path.exists() and not list(tmp_path.glob('.stargate-report-*'))
    finally:release.set()


def test_pdf_ui_export_and_save_dialog_cancellation(demo,tmp_path,wait_until,monkeypatch):
    view,data=demo;page=view.page
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:('','PDF'))
    page.pdf_button.click();assert not page.busy
    path=tmp_path/'viewer.pdf'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:(str(path),'PDF'))
    page.pdf_button.click()
    wait_until(lambda:not page.busy,timeout=15000)
    assert path.read_bytes().startswith(b'%PDF-') and 'Ekspor selesai' in page.message.text()
